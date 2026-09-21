"""
Conciliación de cobros con tarjeta / QR / Mercado Pago - Grupo Oroño

Cruza los recibos de caja (columnas de tarjeta y videoconsultas) contra:
  - Movimientos presentados a Payway (CSV)
  - QR transferencia de Payway (CSV)
  - Movimientos de Mercado Pago (XLSX)

Reglas de conciliación (sin desempates ni tolerancias):
  1. Par directo: misma FECHA y mismo IMPORTE exacto, y esa combinación
     identifica a un único recibo y a un único movimiento.
  2. Grupo que cierra: dentro de una misma fecha e importe, si la cantidad de
     recibos es exactamente igual a la de movimientos, el grupo cierra y se
     concilia en bloque. El apareo dentro del grupo va por orden de carga: el
     total del grupo está verificado, el par individual no.
  3. Diferencia de centavos: sobre lo que quedó sin conciliar, misma FECHA y
     una diferencia de importe de hasta $0,99 en cualquier dirección, siempre
     que sea única de los dos lados (el QR redondea al peso entero).
  4. Neteo de anulaciones: después de las tres anteriores, si entre los recibos
     que quedaron SIN conciliar hay, en la misma FECHA y para el mismo CLIENTE,
     uno positivo y uno negativo del mismo importe, los dos se cancelan entre sí
     y salen del informe (la anulación nunca llegó al procesador). Después se
     vuelven a correr las reglas 1 a 3 sobre lo que quedó. Los recibos negativos
     que YA conciliaron contra una devolución real no se netean nunca.
  5. Pagador propio: los movimientos cuyo campo "Pagador:" es Grupo Oroño los hace
     la empresa, no salen de un recibo de caja. Se apartan ANTES de conciliar y se
     informan aparte. Solo se mira el pagador: "Producto de Grupooroño" aparece en
     cobros de pacientes reales y NO los excluye.
  En las tres primeras, los recibos TVIR (videoconsultas) solo se cruzan contra MP; un
  grupo que empata en cantidad pero no tiene movimientos de MP suficientes para
  sus videoconsultas NO cierra. Todo lo demás queda para revisión manual.
  6. Videoconsulta: un movimiento cuyo detalle dice "Videoconsulta médica" solo
     puede conciliar contra un recibo de una columna virtual (TVIR). Nunca contra
     un recibo de mostrador, aunque la fecha y el importe coincidan.

Uso: dejar los archivos en la carpeta "Entrada" y ejecutar. El resultado queda
en "Salida/Conciliacion_<desde>_<hasta>.xlsx".
"""
import sys
import unicodedata
import warnings
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

COLS_TARJETA = ["TSNS", "TSNI", "TSNZ"]  # amarillas: cobros con tarjeta / QR
COLS_MP = ["TVIR"]  # naranja: videoconsultas, solo Mercado Pago
TOLERANCIA = 0.99  # regla 3: diferencia máxima de importe, en pesos
PAGADOR_PROPIO = "orono"  # regla 5: sin acentos ni mayúsculas

if getattr(sys, "frozen", False):
    BASE = Path(sys.executable).parent
else:
    BASE = Path(__file__).parent
ENTRADA = BASE / "Entrada"
SALIDA = BASE / "Salida"


# ---------------------------------------------------------------- lectura
def leer_csv(path, skiprows=0):
    for enc in ("utf-8-sig", "latin1"):
        try:
            return pd.read_csv(path, sep=";", skiprows=skiprows, encoding=enc, dtype=str)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"No se pudo leer {path.name}")


def num(s):
    return pd.to_numeric(s, errors="coerce")


def leer_recibos(path):
    raw = pd.read_excel(path, header=None)
    fila = raw.index[raw.iloc[:, 0].astype(str).str.strip() == "Cajero"][0]
    df = pd.read_excel(path, header=fila)
    df.columns = [str(c).strip() for c in df.columns]
    faltan = [c for c in COLS_TARJETA + COLS_MP if c not in df.columns]
    if faltan:
        print(f"  ATENCIÓN: {path.name} no tiene las columnas {faltan}")
    filas = []
    for _, x in df.iterrows():
        for c in COLS_TARJETA + COLS_MP:
            if c not in df.columns:
                continue
            v = num(x[c])
            if pd.notna(v) and v != 0:
                filas.append({
                    "Fecha": pd.Timestamp(x["Fecha Caja"]).normalize(),
                    "Nro Recibo": int(x["Nro Recibo"]),
                    "Cliente": str(x["Nombre Cliente"]).strip(),
                    "Cajero": str(x["Cajero"]).strip(),
                    "Columna": c,
                    "Tipo Asiento": str(x["Tipo de Asiento"]).strip(),
                    "Importe": round(float(v), 2),
                    "Importe Total Recibo": round(float(num(x["Importe Total"])), 2),
                })
    return pd.DataFrame(filas)


def leer_payway(path):
    df = leer_csv(path)
    return pd.DataFrame({
        "Fuente": "PAYWAY",
        "Fecha": pd.to_datetime(df["COMPRA"], dayfirst=True),
        "Importe": num(df["MONTO_BRUTO"]).round(2),
        "Referencia": "Lote " + df["LOTE"].str.strip() + " / Cupón " + df["NUM.CUPON"].str.strip()
                      + " / Aut " + df["NRO_AUT"].str.strip(),
        "Detalle": df["TIPO"].str.strip() + " - " + df["MARCA"].str.strip() + " - "
                   + df["DETALLE"].str.strip() + " - Tarj " + df["NUM.TARJETA"].str[-4:],
        "Terminal / Caja": "Estab " + df["ESTABLECIMIENTO"].str.strip(),
    })


def leer_qr(path):
    df = leer_csv(path, skiprows=1)
    return pd.DataFrame({
        "Fuente": "QR",
        "Fecha": pd.to_datetime(df["FECHA DE VENTA"]),
        "Importe": num(df["MONTO BRUTO"]).round(2),
        "Referencia": "QR " + df["QR ID"].str.strip() + " / Cupón " + df["NRO CUPON"].str.strip(),
        "Detalle": df["ESTADO"].str.strip(),
        "Terminal / Caja": "Terminal " + df["TERMINAL LOGICA"].str.strip(),
    })


def leer_mp(path):
    df = pd.read_excel(path, dtype={"ID DE OPERACIÓN EN MERCADO PAGO": str})
    caja = (df["NOMBRE DE LOCAL"].fillna("") + " " + df["NOMBRE DE CAJA"].fillna("")).str.strip()
    detalle = (df["TIPO DE OPERACIÓN"].fillna("") + " - " + df["MEDIO DE PAGO"].fillna("")
               + " - " + df["DETALLE DE LA VENTA"].fillna("").str.replace('"', "")
               + " - Pagador: " + df["PAGADOR"].fillna(""))
    return pd.DataFrame({
        "Fuente": "MP",
        "Fecha": pd.to_datetime(df["FECHA DE ORIGEN"].str[:10]),
        "Importe": num(df["VALOR DE LA COMPRA"]).round(2),
        "Referencia": "MP " + df["ID DE OPERACIÓN EN MERCADO PAGO"].astype(str),
        "Detalle": detalle,
        "Terminal / Caja": caja,
    })


def detectar(path):
    """Identifica cada archivo por su contenido, no por el nombre."""
    ext = path.suffix.lower()
    if ext == ".csv":
        with open(path, encoding="latin1") as f:
            cab = f.read(3000)
        if "Detalle de Transferencias" in cab or "QR ID" in cab:
            return "qr"
        if "NUM.CUPON" in cab:
            return "payway"
    elif ext in (".xls", ".xlsx"):
        cab = pd.read_excel(path, header=None, nrows=8).astype(str).to_string()
        if "Listado de caja" in cab:
            return "recibos"
        if "MERCADO PAGO" in cab.upper():
            return "mp"
    return None


# ---------------------------------------------------------------- conciliación
def conciliar(rec, ext):
    """
    Arma todos los pares posibles recibo-movimiento con misma fecha e importe
    (TVIR solo contra MP). Regla 1: un par se concilia si ninguno de los dos
    tiene otro candidato. Regla 2: si dentro de una fecha + importe la cantidad
    de recibos es igual a la de movimientos, el grupo cierra y se concilia en
    bloque. El resto se agrupa para revisión manual.
    """
    rec = rec.reset_index(drop=True)
    ext = ext.reset_index(drop=True)
    rec["_r"] = rec.index
    ext["_e"] = ext.index

    ext["_video"] = ext.Detalle.map(es_videoconsulta)

    pares = rec.merge(ext, on=["Fecha", "Importe"], suffixes=("", "_mov"))
    pares = pares[~(pares.Columna.isin(COLS_MP) & (pares.Fuente != "MP"))]
    pares = pares[~(~pares.Columna.isin(COLS_MP) & pares._video)]  # regla 6

    cand_r = pares.groupby("_r").size()
    cand_e = pares.groupby("_e").size()
    pares["unico"] = pares._r.map(cand_r).eq(1) & pares._e.map(cand_e).eq(1)

    unicos = pares[pares.unico].copy()
    ambiguos = pares[~pares.unico]

    rec["Estado"] = "Sin movimiento con misma fecha e importe"
    rec.loc[rec._r.isin(ambiguos._r), "Estado"] = "Para revisar"
    rec.loc[rec._r.isin(unicos._r), "Estado"] = "Conciliado"
    ext["Estado"] = "Sin recibo con misma fecha e importe"
    ext.loc[ext._e.isin(ambiguos._e), "Estado"] = "Para revisar"
    ext.loc[ext._e.isin(unicos._e), "Estado"] = "Conciliado"
    unicos["Apareo"] = "Par directo"

    # regla 2: grupos de la misma fecha e importe que empatan en cantidad
    por_ext = ext.groupby(["Fecha", "Importe"])._e.apply(list)
    nuevos = []
    for clave, ridx in rec.groupby(["Fecha", "Importe"])._r.apply(list).items():
        eidx = por_ext.get(clave)
        if eidx is None or len(eidx) != len(ridx):
            continue
        if (rec.loc[ridx, "Estado"] == "Conciliado").any():
            continue  # ya resuelto por la regla 1
        # las videoconsultas solo pueden ir contra MP: tiene que haber al menos
        # tantos movimientos de MP como recibos TVIR
        solo_mp = [i for i in ridx if rec.at[i, "Columna"] in COLS_MP]
        resto = [i for i in ridx if rec.at[i, "Columna"] not in COLS_MP]
        video = [j for j in eidx if ext.at[j, "_video"] and ext.at[j, "Fuente"] == "MP"]
        mp = [j for j in eidx if ext.at[j, "Fuente"] == "MP" and not ext.at[j, "_video"]]
        otros = [j for j in eidx if ext.at[j, "Fuente"] != "MP" and not ext.at[j, "_video"]]
        # una videoconsulta que no sea de MP no aparea con nada: el grupo no cierra
        if len(eidx) != len(video) + len(mp) + len(otros):
            continue
        if len(video) > len(solo_mp):  # regla 6: solo entran en recibos virtuales
            continue
        if len(solo_mp) > len(video) + len(mp):
            continue
        huecos = len(solo_mp) - len(video)
        izq = solo_mp + resto
        der = video + mp[:huecos] + otros + mp[huecos:]
        for i, j in zip(izq, der):
            rec.at[i, "Estado"] = "Conciliado"
            ext.at[j, "Estado"] = "Conciliado"
            nuevos.append((i, j, f"Grupo de {len(izq)}"))

    if nuevos:
        ri = [n[0] for n in nuevos]
        ei = [n[1] for n in nuevos]
        izq_df = rec.loc[ri].drop(columns=["Estado"]).reset_index(drop=True)
        der_df = ext.loc[ei].drop(columns=["Estado", "Fecha", "Importe"]).reset_index(drop=True)
        grupo = pd.concat([izq_df, der_df], axis=1)
        grupo["Apareo"] = [n[2] for n in nuevos]
        unicos = pd.concat([unicos, grupo], ignore_index=True)

    # regla 3: misma fecha y diferencia de hasta $0,99, única de los dos lados,
    # sobre lo que quedó sin conciliar por las reglas 1 y 2
    libre_r = rec[rec.Estado != "Conciliado"]
    libre_e = ext[ext.Estado != "Conciliado"]
    cerca = libre_r.merge(libre_e, on="Fecha", suffixes=("", "_mov"))
    cerca = cerca[(cerca.Importe - cerca.Importe_mov).abs().round(2) <= TOLERANCIA]
    cerca = cerca[~(cerca.Columna.isin(COLS_MP) & (cerca.Fuente != "MP"))]
    cerca = cerca[~(~cerca.Columna.isin(COLS_MP) & cerca._video)]  # regla 6
    n_r = cerca.groupby("_r").size()
    n_e = cerca.groupby("_e").size()
    cerca = cerca[cerca._r.map(n_r).eq(1) & cerca._e.map(n_e).eq(1)]
    if len(cerca):
        rec.loc[rec._r.isin(cerca._r), "Estado"] = "Conciliado"
        ext.loc[ext._e.isin(cerca._e), "Estado"] = "Conciliado"
        cerca = cerca.drop(columns=["Estado", "Estado_mov"])
        cerca["Apareo"] = "Diferencia hasta $0,99"
        unicos = pd.concat([unicos, cerca], ignore_index=True)
    unicos["Importe mov"] = unicos.Importe_mov.fillna(unicos.Importe) if "Importe_mov" in unicos else unicos.Importe
    unicos["Diferencia"] = (unicos["Importe mov"] - unicos.Importe).round(2)
    return rec, ext, unicos


def sin_acentos(s):
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()


def es_videoconsulta(detalle):
    """Regla 6: ¿el movimiento es una videoconsulta? Solo va contra un recibo virtual."""
    return "videoconsulta" in sin_acentos(detalle)


def pagador_propio(detalle):
    """Regla 5: ¿el movimiento lo pagó la propia empresa? Solo mira el campo "Pagador:"."""
    d = sin_acentos(detalle)
    i = d.find("pagador:")
    return i >= 0 and PAGADOR_PROPIO in d[i + len("pagador:"):]


def netear(rec):
    """
    Regla 4: entre los recibos SIN conciliar, cancela los pares positivo/negativo
    de la misma fecha, el mismo cliente y el mismo importe. Devuelve el DataFrame
    de pares anulados (una fila por par) y la lista de _r que salen del informe.
    """
    libre = rec[(rec.Estado != "Conciliado") & (rec.Importe != 0)].copy()
    libre["_abs"] = libre.Importe.abs().round(2)
    pares, fuera = [], []
    for _, g in libre.groupby(["Fecha", "Cliente", "_abs"]):
        pos = g[g.Importe > 0].sort_values("Nro Recibo")
        neg = g[g.Importe < 0].sort_values("Nro Recibo")
        n = min(len(pos), len(neg))
        if not n:
            continue
        pos, neg = pos.head(n), neg.head(n)
        fuera += list(pos._r) + list(neg._r)
        for (_, p), (_, q) in zip(pos.iterrows(), neg.iterrows()):
            pares.append({"Fecha": p.Fecha, "Cliente": p.Cliente, "Importe": p.Importe,
                          "Recibo": p["Nro Recibo"], "Cajero": p.Cajero, "Columna": p.Columna,
                          "Tipo Asiento": p["Tipo Asiento"], "Recibo anulación": q["Nro Recibo"],
                          "Cajero anulación": q.Cajero, "Columna anulación": q.Columna,
                          "Tipo Asiento anulación": q["Tipo Asiento"], "Importe anulación": q.Importe})
    cols = ["Fecha", "Cliente", "Importe", "Recibo", "Cajero", "Columna", "Tipo Asiento",
            "Recibo anulación", "Cajero anulación", "Columna anulación",
            "Tipo Asiento anulación", "Importe anulación"]
    anulados = pd.DataFrame(pares, columns=cols)
    if len(anulados):
        anulados = anulados.sort_values(["Fecha", "Recibo"])
    return anulados, fuera


# ---------------------------------------------------------------- salida
COLS_REC = ["Fecha", "Nro Recibo", "Cliente", "Cajero", "Columna", "Tipo Asiento", "Importe"]
COLS_EXT = ["Fuente", "Fecha", "Importe", "Referencia", "Detalle", "Terminal / Caja"]
COLS_REV = ["Grupo", "Fecha", "Importe recibo", "Importe movimiento", "Cant. recibos",
            "Cant. movimientos", "Lado", "Nro Recibo", "Cliente", "Cajero", "Columna",
            "Tipo Asiento", "Fuente", "Referencia", "Detalle", "Terminal / Caja"]


def hoja_grupos(r, e, con_estado=False):
    """Por cada fecha + importe: todos los recibos y movimientos juntos, uno debajo del otro."""
    filas = []
    grupos = sorted(set(zip(r.Fecha, r.Importe)) | set(zip(e.Fecha, e.Importe)))
    for n, (f, imp) in enumerate(grupos, 1):
        gr = r[(r.Fecha == f) & (r.Importe == imp)]
        ge = e[(e.Fecha == f) & (e.Importe == imp)]
        # el importe va en la columna del lado que corresponde, para sumar cada uno por separado
        base = {"Grupo": n, "Fecha": f, "Cant. recibos": len(gr), "Cant. movimientos": len(ge)}
        for _, x in gr.iterrows():
            est = {"Estado": x.Estado if x.Estado == "Para revisar" else "Recibo sin movimiento"} if con_estado else {}
            filas.append({**est, **base, "Importe recibo": imp,
                          "Lado": "Recibo", "Nro Recibo": x["Nro Recibo"],
                          "Cliente": x.Cliente, "Cajero": x.Cajero, "Columna": x.Columna,
                          "Tipo Asiento": x["Tipo Asiento"]})
        for _, x in ge.iterrows():
            est = {"Estado": x.Estado if x.Estado == "Para revisar" else "Movimiento sin recibo"} if con_estado else {}
            filas.append({**est, **base, "Importe movimiento": imp,
                          "Lado": "Movimiento", "Fuente": x.Fuente,
                          "Referencia": x.Referencia, "Detalle": x.Detalle,
                          "Terminal / Caja": x["Terminal / Caja"]})
    cols = (["Estado"] if con_estado else []) + COLS_REV
    df = pd.DataFrame(filas).reindex(columns=cols)
    df["Nro Recibo"] = df["Nro Recibo"].astype("Int64")
    return df


def armar_salida(rec, ext, unicos, anulados, propios, desde, hasta, archivos):
    conc = unicos.rename(columns={"Referencia": "Referencia mov", "Detalle": "Detalle mov"})
    conc = conc[["Fecha", "Importe", "Apareo", "Nro Recibo", "Cliente", "Cajero", "Columna",
                 "Tipo Asiento", "Fuente", "Importe mov", "Diferencia", "Referencia mov",
                 "Detalle mov", "Terminal / Caja"]]
    conc = conc.sort_values(["Fecha", "Nro Recibo"])

    revisar = hoja_grupos(rec[rec.Estado == "Para revisar"], ext[ext.Estado == "Para revisar"])
    trabajo = hoja_grupos(rec[rec.Estado != "Conciliado"], ext[ext.Estado != "Conciliado"], True)
    rec_sin = rec[rec.Estado.str.startswith("Sin")][COLS_REC].sort_values(["Fecha", "Nro Recibo"])
    ext_sin = ext[ext.Estado.str.startswith("Sin")][COLS_EXT].sort_values(["Fuente", "Fecha"])

    res = [["Período", f"{desde:%d/%m/%Y} al {hasta:%d/%m/%Y}", "", ""],
           ["Procesado", datetime.now().strftime("%d/%m/%Y %H:%M"), "", ""],
           ["Regla 1", "Par directo: misma fecha + mismo importe exacto, único de ambos lados", "", ""],
           ["Regla 2", "Grupo que cierra: misma fecha + mismo importe, igual cantidad de recibos que de movimientos", "", ""],
           ["Regla 3", "Diferencia de centavos: misma fecha + diferencia de hasta $0,99, única de ambos lados", "", ""],
           ["Regla 4", "Neteo de anulaciones: recibo positivo y negativo sin conciliar, misma fecha, mismo cliente e importe", "", ""],
           ["Regla 5", "Pagador Grupo Oroño: los movimientos que paga la empresa se apartan antes de conciliar", "", ""],
           ["Regla 6", 'Videoconsulta: un movimiento que dice "Videoconsulta médica" solo concilia contra un recibo ' + ", ".join(COLS_MP), "", ""],
           ["", "En las tres primeras, TVIR solo se cruza contra Mercado Pago", "", ""],
           ["", "", "", ""],
           ["RECIBOS", "Líneas", "Importe", ""]]
    for col in COLS_TARJETA + COLS_MP:
        g = rec[rec.Columna == col]
        if len(g):
            res.append([col, len(g), round(g.Importe.sum(), 2), ""])
    res.append(["Total", len(rec), round(rec.Importe.sum(), 2), ""])
    res.append(["", "", "", ""])
    res.append(["MOVIMIENTOS", "Cantidad", "Importe", ""])
    for f in ["PAYWAY", "QR", "MP"]:
        g = ext[ext.Fuente == f]
        if len(g):
            res.append([f, len(g), round(g.Importe.sum(), 2), ""])
    res.append(["Total", len(ext), round(ext.Importe.sum(), 2), ""])
    res.append(["", "", "", ""])
    res.append(["RESULTADO", "Recibos", "Importe recibos", "Movimientos"])
    for estado in ["Conciliado", "Para revisar"]:
        g, h = rec[rec.Estado == estado], ext[ext.Estado == estado]
        res.append([estado, len(g), round(g.Importe.sum(), 2), len(h)])
        if estado == "Conciliado":
            for etiqueta, sel in [("   por par directo (regla 1)", unicos.Apareo == "Par directo"),
                                  ("   por grupo que cierra (regla 2)", unicos.Apareo.str.startswith("Grupo")),
                                  ("   por diferencia de hasta $0,99 (regla 3)",
                                   unicos.Apareo == "Diferencia hasta $0,99")]:
                g = unicos[sel]
                res.append([etiqueta, len(g), round(g.Importe.sum(), 2), len(g)])
    g = rec[rec.Estado.str.startswith("Sin")]
    res.append(["Recibos sin movimiento", len(g), round(g.Importe.sum(), 2), ""])
    h = ext[ext.Estado.str.startswith("Sin")]
    res.append(["Movimientos sin recibo", "", round(h.Importe.sum(), 2), len(h)])
    res.append(["Anulados y neteados (regla 4)", len(anulados) * 2, 0, ""])
    res.append(["Pagador Grupo Oroño (estos no tienen recibo)", "",
                round(propios.Importe.sum(), 2) if len(propios) else 0, len(propios)])
    res.append(["", "", "", ""])
    res.append(["ARCHIVOS LEÍDOS", "", "", ""])
    for tipo, nombre in archivos:
        res.append([tipo, nombre, "", ""])
    resumen = pd.DataFrame(res)

    hojas = {"Resumen": resumen, "Conciliados": conc, "Anulados": anulados}
    if len(propios):
        hojas["Pagador Grupo Oroño"] = propios[COLS_EXT].sort_values(["Fuente", "Fecha"])
    hojas.update({"Hoja de trabajo": trabajo, "Para revisar": revisar,
                  "Recibos sin movimiento": rec_sin, "Movimientos sin recibo": ext_sin})
    return hojas


def escribir_excel(hojas, destino):
    with pd.ExcelWriter(destino, engine="openpyxl") as w:
        for nombre, df in hojas.items():
            df.to_excel(w, sheet_name=nombre, index=False, header=(nombre != "Resumen"))
    wb = load_workbook(destino)
    azul = PatternFill("solid", fgColor="1F4E78")
    gris = PatternFill("solid", fgColor="EDEDED")
    for ws in wb.worksheets:
        es_resumen = ws.title == "Resumen"
        encabezados = [c.value for c in ws[1]]
        if not es_resumen:
            for cel in ws[1]:
                cel.font = Font(bold=True, color="FFFFFF")
                cel.fill = azul
                cel.alignment = Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letra = get_column_letter(col[0].column)
            titulo = str(col[0].value or "")
            largo = max(len(str(c.value)) if c.value is not None else 0 for c in col[:300])
            ws.column_dimensions[letra].width = min(max(largo, len(titulo), 8) + 2, 60)
            for c in (col if es_resumen else col[1:]):
                if isinstance(c.value, datetime):
                    c.number_format = "DD/MM/YYYY"
                elif isinstance(c.value, float):
                    c.number_format = "#,##0.00"
        if es_resumen:
            for row in ws.iter_rows():
                if row[0].value and str(row[0].value).isupper():
                    for c in row:
                        c.font = Font(bold=True)
        if "Grupo" in encabezados:
            # sombreado alternado por grupo para leerlo más fácil
            g = encabezados.index("Grupo")
            for row in ws.iter_rows(min_row=2):
                if row[g].value and row[g].value % 2 == 0:
                    for c in row:
                        c.fill = gris
    wb.save(destino)


# ---------------------------------------------------------------- main
def main():
    ENTRADA.mkdir(exist_ok=True)
    SALIDA.mkdir(exist_ok=True)
    lectores = {"recibos": leer_recibos, "payway": leer_payway, "qr": leer_qr, "mp": leer_mp}
    datos = {k: [] for k in lectores}
    archivos = []

    print(f"Leyendo archivos de {ENTRADA}")
    for p in sorted(ENTRADA.iterdir()):
        if p.name.startswith("~$") or not p.is_file():
            continue
        tipo = detectar(p)
        if tipo is None:
            print(f"  - {p.name}: no reconocido, se ignora")
            continue
        print(f"  - {p.name}: {tipo}")
        archivos.append((tipo, p.name))
        datos[tipo].append(lectores[tipo](p))

    for req in lectores:
        if not datos[req]:
            print(f"  ATENCIÓN: falta el archivo de {req}")
    if not datos["recibos"]:
        sys.exit("No hay archivo de recibos: no se puede conciliar.")

    rec = pd.concat(datos["recibos"]).drop_duplicates()
    ext = pd.concat([d for k in ("payway", "qr", "mp") for d in datos[k]]).drop_duplicates()
    desde, hasta = rec.Fecha.min(), rec.Fecha.max()

    # regla 5: los movimientos pagados por la propia empresa no se concilian
    propios = ext[ext.Detalle.map(pagador_propio)].copy()
    if len(propios):
        print(f"  aparta {len(propios)} movimientos con pagador Grupo Oroño (regla 5)")
        ext = ext[~ext.Detalle.map(pagador_propio)]

    print(f"\nConciliando {len(rec)} líneas de recibos contra {len(ext)} movimientos...")
    rec, ext, unicos = conciliar(rec, ext)
    anulados, fuera = netear(rec)
    if fuera:
        print(f"  netea {len(fuera)} recibos en {len(anulados)} pares anulados (regla 4)")
        rec = rec[~rec._r.isin(fuera)].drop(columns=["_r", "Estado"])
        ext = ext.drop(columns=["_e", "Estado"])
        rec, ext, unicos = conciliar(rec, ext)
    hojas = armar_salida(rec, ext, unicos, anulados, propios, desde, hasta, archivos)

    destino = SALIDA / f"Conciliacion_{desde:%Y-%m-%d}_{hasta:%Y-%m-%d}.xlsx"
    escribir_excel(hojas, destino)

    for estado in ["Conciliado", "Para revisar"]:
        print(f"{estado:<24} recibos: {(rec.Estado == estado).sum():>4}   "
              f"movimientos: {(ext.Estado == estado).sum():>4}")
    print(f"{'Recibos sin movimiento':<24} {len(hojas['Recibos sin movimiento']):>13}")
    print(f"{'Movimientos sin recibo':<24} {len(hojas['Movimientos sin recibo']):>32}")
    print(f"\nResultado: {destino}")


if __name__ == "__main__":
    try:
        main()
    except SystemExit as e:
        print(e)
    finally:
        if getattr(sys, "frozen", False):
            input("\nPresioná Enter para cerrar...")
